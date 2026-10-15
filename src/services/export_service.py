"""Excel and CSV exports of reconciliation results.

Security: text from client documents is written as text and values that start
with ``= + - @`` are prefixed with an apostrophe so spreadsheet applications
never evaluate them as formulas (CSV/formula injection).
"""

from __future__ import annotations

import csv
import io
from collections.abc import Callable, Sequence
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from typing import Any, cast

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from src.database.models import ReconciliationRow
from src.models.reconciliation import MatchStatus, ResultKind
from src.services.reconciliation_service import StoredRun
from src.services.result_views import ResultRow, summary_messages

MONEY_FORMAT = "#,##0.00;[Red]-#,##0.00"
DATE_FORMAT = "dd-mmm-yyyy"
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14, color="1F4E79")
LABEL_FONT = Font(bold=True, color="5B6675")
THIN = Side(style="thin", color="E2E7EF")
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def safe_text(value: object) -> object:
    """Neutralise spreadsheet formula injection for text values."""
    if isinstance(value, str) and value.startswith(_FORMULA_PREFIXES):
        return "'" + value
    return value


@dataclass
class ExportContext:
    recon: ReconciliationRow
    run: StoredRun
    rows: list[ResultRow]
    audit: list[dict] = field(default_factory=list)
    generated_at: datetime = field(default_factory=datetime.now)

    @property
    def supplier(self) -> str:
        return self.recon.supplier.name


def _status(*statuses: MatchStatus) -> Callable[[ResultRow], bool]:
    wanted = set(statuses)
    return lambda row: row.status in wanted


SHEETS: tuple[tuple[str, Callable[[ResultRow], bool]], ...] = (
    ("Matched", _status(MatchStatus.MATCHED, MatchStatus.MATCHED_WITH_TOLERANCE,
                        MatchStatus.CREDIT_NOTE)),
    ("Missing", _status(MatchStatus.MISSING_IN_ACCOUNTING)),
    ("Duplicates", _status(MatchStatus.DUPLICATE)),
    ("Amount Mismatches", _status(MatchStatus.AMOUNT_MISMATCH)),
    ("Date Mismatches", _status(MatchStatus.DATE_MISMATCH)),
    ("Reference Mismatches", _status(MatchStatus.REFERENCE_MISMATCH)),
    ("Credit Notes", lambda row: row.kind == ResultKind.CREDIT_NOTE),
    ("Payment Issues", _status(MatchStatus.PAYMENT_MISMATCH)),
    ("Accounting Only", _status(MatchStatus.NOT_ON_SUPPLIER_STATEMENT)),
    ("Unresolved", _status(MatchStatus.UNRESOLVED)),
)  # fmt: skip


def _write_cell(sheet: Worksheet, row: int, column: int, value: object, header: str = ""):
    cell = sheet.cell(row=row, column=column, value=cast(Any, safe_text(value)))
    lower = header.lower()
    if isinstance(value, Decimal) or (
        isinstance(value, int | float) and any(w in lower for w in ("amount", "difference"))
    ):
        cell.number_format = MONEY_FORMAT
    elif isinstance(value, date | datetime):
        cell.number_format = DATE_FORMAT
    return cell


def _write_table(sheet: Worksheet, records: Sequence[dict], start_row: int = 1) -> None:
    if not records:
        sheet.cell(row=start_row, column=1, value="No items in this category.").font = LABEL_FONT
        return
    headers = list(records[0].keys())
    for col, header in enumerate(headers, start=1):
        cell = sheet.cell(row=start_row, column=col, value=header)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for r_index, record in enumerate(records, start=start_row + 1):
        for col, header in enumerate(headers, start=1):
            cell = _write_cell(sheet, r_index, col, record[header], header)
            cell.border = Border(bottom=THIN)
            if header in ("Reason", "Description", "Details", "Data"):
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.freeze_panes = f"A{start_row + 1}"
    sheet.auto_filter.ref = (
        f"A{start_row}:{get_column_letter(len(headers))}{start_row + len(records)}"
    )
    _autosize(sheet)


def _autosize(sheet: Worksheet, minimum: int = 8, maximum: int = 60) -> None:
    widths: dict[int, int] = {}
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            value = cell.value
            if isinstance(value, Decimal | float):
                text = f"{value:,.2f}"
            elif isinstance(value, date | datetime):
                text = "00-Mmm-0000"
            else:
                text = str(value)
            widths[cell.column] = max(widths.get(cell.column, 0), len(text) + 2)
    for column, width in widths.items():
        sheet.column_dimensions[get_column_letter(column)].width = max(minimum, min(maximum, width))


def _key_values(sheet: Worksheet, items: Sequence[tuple[str, object]], start_row: int) -> int:
    row = start_row
    for label, value in items:
        sheet.cell(row=row, column=1, value=label).font = LABEL_FONT
        _write_cell(sheet, row, 2, value, "amount" if isinstance(value, Decimal) else "")
        row += 1
    return row


def summary_items(context: ExportContext) -> list[tuple[str, object]]:
    recon, summary = context.recon, context.run.summary
    period = " to ".join(
        d.strftime("%d-%b-%Y") for d in (recon.period_start, recon.period_end) if d
    )
    outcome = (recon.outcome or summary.suggested_outcome.value).replace("_", " ").title()
    return [
        ("Client / company", recon.client_name or ""),
        ("Supplier", context.supplier),
        ("Supplier account", recon.supplier_account or ""),
        ("Statement period", period),
        ("Currency", recon.currency),
        ("Generated", context.generated_at.strftime("%d-%b-%Y %H:%M")),
        ("Overall status", outcome + ("" if recon.outcome else " (suggested)")),
        ("Total statement transactions", summary.statement_transactions),
        ("Total accounting transactions", summary.accounting_transactions),
        ("Matched", summary.matched),
        ("Missing in accounting", summary.count(MatchStatus.MISSING_IN_ACCOUNTING)),
        ("Not on supplier statement", summary.count(MatchStatus.NOT_ON_SUPPLIER_STATEMENT)),
        ("Duplicates", summary.count(MatchStatus.DUPLICATE)),
        ("Amount mismatches", summary.count(MatchStatus.AMOUNT_MISMATCH)),
        ("Date mismatches", summary.count(MatchStatus.DATE_MISMATCH)),
        ("Reference mismatches", summary.count(MatchStatus.REFERENCE_MISMATCH)),
        ("Payment issues", summary.count(MatchStatus.PAYMENT_MISMATCH)),
        ("Unresolved", summary.count(MatchStatus.UNRESOLVED)),
        ("Pending reviews", summary.pending_reviews),
    ]


def balance_items(context: ExportContext) -> list[tuple[str, object]]:
    balance = context.run.balance
    items: list[tuple[str, object]] = [
        ("Status", balance.status.value.replace("_", " ").title()),
        ("Explanation", balance.reason),
        ("Opening balance", balance.opening_balance),
        ("Invoices", balance.total_invoices),
        ("Credit notes", balance.total_credit_notes),
        ("Payments", balance.total_payments),
        ("Expected closing balance", balance.expected_closing_balance),
        ("Supplier closing balance", balance.supplier_closing_balance),
        ("Difference", balance.difference),
        ("Accounting balance", balance.accounting_balance),
        ("Accounting balance basis", balance.accounting_balance_basis or ""),
        ("Statement vs accounting difference", balance.statement_vs_accounting_difference),
    ]
    items += [(f"  of which: {k}", v) for k, v in balance.explained_difference.items()]
    items.append(("Unexplained difference", balance.unexplained_difference))
    return items


class ExportService:
    def workbook_bytes(self, context: ExportContext) -> bytes:
        workbook = Workbook()
        sheet = workbook.active
        assert isinstance(sheet, Worksheet)
        sheet.title = "Summary"
        sheet["A1"] = "Supplier Statement Reconciliation"
        sheet["A1"].font = TITLE_FONT
        if context.recon.is_demo:
            sheet["A2"] = "DEMO DATA - synthetic records for demonstration only"
            sheet["A2"].font = Font(bold=True, color="9A6200")
        row = _key_values(sheet, summary_items(context), 4) + 1
        sheet.cell(row=row, column=1, value="Key findings").font = LABEL_FONT
        for offset, line in enumerate(
            summary_messages(context.run.summary.status_counts, context.run.summary.pending_reviews)
        ):
            sheet.cell(row=row + 1 + offset, column=1, value=line)
        sheet.column_dimensions["A"].width = 34
        sheet.column_dimensions["B"].width = 48

        for title, predicate in SHEETS:
            _write_table(
                workbook.create_sheet(title),
                [r.as_record() for r in context.rows if predicate(r)],
            )

        balance_sheet = workbook.create_sheet("Balance Check")
        balance_sheet["A1"] = "Statement balance reconciliation"
        balance_sheet["A1"].font = TITLE_FONT
        next_row = _key_values(balance_sheet, balance_items(context), 3)
        breaks = context.run.balance.running_balance_breaks
        if breaks:
            balance_sheet.cell(row=next_row + 1, column=1,
                               value="Running balance breaks").font = LABEL_FONT  # fmt: skip
            for offset, line in enumerate(breaks):
                balance_sheet.cell(row=next_row + 2 + offset, column=1, value=safe_text(line))
        if context.run.allocations:
            start = next_row + len(breaks) + 4
            balance_sheet.cell(
                row=start - 1, column=1, value="Payment allocation"
            ).font = LABEL_FONT
            _write_table(
                balance_sheet,
                [
                    {"Source": a.source.title(), "Invoice": a.invoice_number,
                     "Invoice Amount": a.invoice_amount, "Payments": a.payments_allocated,
                     "Credits": a.credits_allocated, "Outstanding Amount": a.outstanding,
                     "Allocation Status": a.allocation_status.replace("_", " ").title()}
                    for a in context.run.allocations
                ],
                start_row=start,
            )  # fmt: skip
            balance_sheet.freeze_panes = None
        balance_sheet.column_dimensions["A"].width = 40
        balance_sheet.column_dimensions["B"].width = 60

        if context.audit:
            _write_table(workbook.create_sheet("Audit Trail"), context.audit)

        buffer = io.BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    def csv_bytes(self, rows: Sequence[ResultRow]) -> bytes:
        """All results in one CSV (UTF-8 with BOM so Excel detects the encoding)."""
        buffer = io.StringIO()
        records = [row.as_record() for row in rows]
        headers = list(records[0].keys()) if records else list(ResultRow.__annotations__)
        writer = csv.DictWriter(buffer, fieldnames=headers)
        writer.writeheader()
        for record in records:
            writer.writerow({k: _csv_value(v) for k, v in record.items()})
        return buffer.getvalue().encode("utf-8-sig")


def _csv_value(value: object) -> object:
    if value is None:
        return ""
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, Decimal):
        return f"{value:.2f}"
    return safe_text(value)


def build_export_context(service, reconciliation_id: int) -> ExportContext | None:
    """Assemble everything an export needs, or None if the reconciliation has not run."""
    from src.services.result_views import build_rows

    recon = service.get(reconciliation_id)
    run = service.stored_run(reconciliation_id)
    if recon is None or run is None:
        return None
    data = service.load_data(reconciliation_id)
    audit = [
        {"When": e["When"], "Event": e["Event"], "Details": e["Details"]}
        for e in service.audit_trail(reconciliation_id)
    ]
    return ExportContext(recon=recon, run=run, rows=build_rows(run.results, data.lookup()),
                         audit=audit)  # fmt: skip


def export_filename(context: ExportContext, extension: str) -> str:
    import re

    supplier = re.sub(r"[^A-Za-z0-9]+", "_", context.supplier).strip("_")[:40] or "supplier"
    stamp = context.generated_at.strftime("%Y%m%d")
    return f"reconciliation_{context.recon.id}_{supplier}_{stamp}.{extension}"
