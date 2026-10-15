"""Low-level readers that turn CSV / XLSX files into grids of strings.

Shared by statement parsers and accounting providers.
"""

from __future__ import annotations

import csv
import io
import logging
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from src.parsers.columns import header_row_score

logger = logging.getLogger(__name__)

_ENCODINGS = ("utf-8-sig", "cp1252", "latin-1")
_MAX_ROWS = 200_000


class FileReadError(Exception):
    """Raised with a user-facing message when a file cannot be read."""


def _decode(data: bytes) -> str:
    if b"\x00" in data[:4096]:
        raise FileReadError(
            "This file does not look like a text CSV file. If it is a spreadsheet, "
            "save it as .xlsx or .csv and upload it again."
        )
    for encoding in _ENCODINGS:
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise FileReadError("The file's text encoding could not be recognised.")


def read_csv_grid(path: Path) -> list[list[str]]:
    """Read a CSV file with automatic encoding and delimiter detection."""
    data = path.read_bytes()
    if not data.strip():
        raise FileReadError("The file is empty.")
    text = _decode(data)
    sample = text[:20_000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ","
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    grid = []
    for index, row in enumerate(reader):
        if index >= _MAX_ROWS:
            raise FileReadError(f"The file has more than {_MAX_ROWS:,} rows.")
        grid.append(row)
    if not any(any(cell.strip() for cell in row) for row in grid):
        raise FileReadError("The file contains no data.")
    return grid


def excel_cell_to_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time(0) else value.isoformat(" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float):
        return repr(value) if not value.is_integer() else str(int(value))
    if isinstance(value, Decimal | int):
        return str(value)
    return str(value)


def list_excel_sheets(path: Path) -> list[str]:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises many types for corrupt files
        logger.warning("Could not open workbook: %s", type(exc).__name__)
        raise FileReadError(
            "The Excel file could not be opened. Please check that it is a valid .xlsx file "
            "(older .xls files must be re-saved as .xlsx)."
        ) from exc
    try:
        return [ws.title for ws in workbook.worksheets if ws.sheet_state == "visible"]
    finally:
        workbook.close()


def read_excel_grid(path: Path, sheet_name: str | None = None) -> tuple[list[list[str]], str]:
    """Read one worksheet. Without ``sheet_name`` the most table-like sheet is chosen."""
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        logger.warning("Could not open workbook: %s", type(exc).__name__)
        raise FileReadError(
            "The Excel file could not be opened. Please check that it is a valid .xlsx file "
            "(older .xls files must be re-saved as .xlsx)."
        ) from exc

    try:
        sheets = [ws for ws in workbook.worksheets if ws.sheet_state == "visible"]
        if sheet_name:
            sheets = [ws for ws in sheets if ws.title == sheet_name]
            if not sheets:
                raise FileReadError(f"Worksheet '{sheet_name}' was not found.")

        best: tuple[int, int, list[list[str]], str] | None = None
        for sheet in sheets:
            grid: list[list[str]] = []
            for index, row in enumerate(sheet.iter_rows(values_only=True)):
                if index >= _MAX_ROWS:
                    raise FileReadError(f"Worksheet '{sheet.title}' has too many rows.")
                grid.append([excel_cell_to_text(v) for v in row])
            while grid and not any(c.strip() for c in grid[-1]):
                grid.pop()
            header_score = max((header_row_score(r) for r in grid[:40]), default=0)
            candidate = (header_score, len(grid), grid, sheet.title)
            if best is None or candidate[:2] > best[:2]:
                best = candidate
    finally:
        workbook.close()

    if best is None or not best[2]:
        raise FileReadError("The workbook contains no data.")
    return best[2], best[3]
