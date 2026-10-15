"""Excel/CSV exports and the PDF client report."""

from __future__ import annotations

import csv
import io
from dataclasses import replace
from decimal import Decimal

import pdfplumber
import pytest
from openpyxl import load_workbook

from src.config import get_settings
from src.services.export_service import (
    ExportService,
    build_export_context,
    export_filename,
    safe_text,
)
from src.services.import_service import FileRole, ImportService
from src.services.reconciliation_service import NewReconciliation, ReconciliationService
from src.services.report_service import ReportService

STATEMENT = b"""Date,Invoice No,Details,Debit,Credit,Balance
01/09/2026,,Opening balance,,,100.00
03/09/2026,INV-1001,Widgets,1200.00,,1300.00
05/09/2026,INV-1002,=HYPERLINK("http://evil"),300.00,,1600.00
06/09/2026,CN-1,Credit note,,50.00,1550.00
30/09/2026,,Closing balance,,,1550.00
"""
BILLS = b"""Bill No,Bill Date,Total,Type
INV-1001,03/09/2026,1200.00,Bill
INV-1002,05/09/2026,280.00,Bill
INV-1007,07/09/2026,90.00,Bill
"""

EXPECTED_SHEETS = [
    "Summary", "Matched", "Missing", "Duplicates", "Amount Mismatches", "Date Mismatches",
    "Reference Mismatches", "Credit Notes", "Payment Issues", "Accounting Only", "Unresolved",
    "Balance Check", "Audit Trail",
]  # fmt: skip


@pytest.fixture
def context(database, tmp_path):
    service = ReconciliationService(database)
    importer = ImportService(database, replace(get_settings(), upload_dir=tmp_path / "up"))
    recon_id = service.create(
        NewReconciliation(supplier_name="Acme Supplies Ltd", currency="GBP", client_name="Client")
    )
    importer.upload(recon_id, FileRole.STATEMENT, "s.csv", STATEMENT)
    importer.upload(recon_id, FileRole.ACCOUNTING, "b.csv", BILLS)
    service.run(recon_id)
    return build_export_context(service, recon_id)


def test_workbook_has_required_sheets_and_formatting(context):
    workbook = load_workbook(io.BytesIO(ExportService().workbook_bytes(context)))
    assert workbook.sheetnames == EXPECTED_SHEETS

    amounts = workbook["Amount Mismatches"]
    assert amounts.freeze_panes == "A2"
    headers = [c.value for c in amounts[1]]
    row = dict(zip(headers, [c.value for c in amounts[2]], strict=True))
    assert row["Reference"] == "INV-1002"
    assert row["Difference"] == 20
    difference_cell = amounts.cell(row=2, column=headers.index("Difference") + 1)
    assert "#,##0.00" in difference_cell.number_format
    date_cell = amounts.cell(row=2, column=headers.index("Statement Date") + 1)
    assert date_cell.number_format == "dd-mmm-yyyy"
    assert amounts.column_dimensions["A"].width >= 8

    summary = {r[0].value: r[1].value for r in workbook["Summary"].iter_rows(min_row=4)
               if r[0].value}  # fmt: skip
    assert summary["Supplier"] == "Acme Supplies Ltd"
    assert summary["Amount mismatches"] == 1
    assert summary["Missing in accounting"] == 1  # credit note CN-1
    credit_sheet = workbook["Credit Notes"]
    assert credit_sheet.cell(row=2, column=3).value == "CN-1"
    assert workbook["Accounting Only"].cell(row=2, column=3).value == "INV-1007"
    assert workbook["Duplicates"]["A1"].value == "No items in this category."


def test_formula_injection_is_neutralised(context):
    workbook = load_workbook(io.BytesIO(ExportService().workbook_bytes(context)))
    values = [c.value for row in workbook["Amount Mismatches"].iter_rows() for c in row]
    assert any(isinstance(v, str) and v.startswith("'=HYPERLINK") for v in values)
    assert not any(isinstance(v, str) and v.startswith("=") for v in values)
    assert safe_text("-1") == "'-1"
    assert safe_text("INV-1") == "INV-1"
    assert safe_text(Decimal("-1")) == Decimal("-1")


def test_csv_export(context):
    data = ExportService().csv_bytes(context.rows)
    assert data.startswith(b"\xef\xbb\xbf")  # BOM for Excel
    rows = list(csv.DictReader(io.StringIO(data.decode("utf-8-sig"))))
    assert len(rows) == len(context.rows)
    mismatch = next(r for r in rows if r["Reference"] == "INV-1002")
    assert mismatch["Status"] == "Amount mismatch"
    assert mismatch["Difference"] == "20.00"
    assert mismatch["Statement Date"] == "2026-09-05"


def test_pdf_report(context, tmp_path):
    pdf = ReportService().pdf_bytes(context)
    assert pdf.startswith(b"%PDF")
    path = tmp_path / "report.pdf"
    path.write_bytes(pdf)
    with pdfplumber.open(path) as document:
        text = " ".join(page.extract_text() or "" for page in document.pages)
    assert "Supplier Statement Reconciliation" in text
    assert "Acme Supplies Ltd" in text
    assert "Balance reconciliation" in text
    assert "Amount differences (1)" in text
    assert "INV-1007" in text
    assert "Overall status" in text


def test_export_filename(context):
    assert export_filename(context, "xlsx").startswith(f"reconciliation_{context.recon.id}_Acme")
