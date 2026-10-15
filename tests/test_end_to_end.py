"""End-to-end: demo files -> import -> normalise -> reconcile -> review -> export."""

from __future__ import annotations

import io
from dataclasses import replace
from decimal import Decimal

import pytest
import streamlit as st
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

from src import config
from src.config import PROJECT_ROOT, get_settings
from src.database import database as database_module
from src.models.reconciliation import (
    BalanceStatus,
    MatchMethod,
    MatchStatus,
    ReconciliationOutcome,
    ReviewStatus,
)
from src.services import demo_data
from src.services.export_service import ExportService, build_export_context
from src.services.import_service import FileRole, ImportService
from src.services.reconciliation_service import NewReconciliation, ReconciliationService
from src.services.report_service import ReportService
from src.services.result_views import build_rows
from src.services.review_service import ReviewAction, ReviewService
from ui.state import CURRENT_KEY

EXPECTED_STATUS = {
    "exact": MatchStatus.MATCHED,
    "reference_format": MatchStatus.MATCHED,
    "within_tolerance": MatchStatus.MATCHED_WITH_TOLERANCE,
    "missing_in_accounting": MatchStatus.MISSING_IN_ACCOUNTING,
    "amount_mismatch": MatchStatus.AMOUNT_MISMATCH,
    "date_mismatch": MatchStatus.DATE_MISMATCH,
    "fuzzy_typo": MatchStatus.REFERENCE_MISMATCH,
    "credit_note_matched": MatchStatus.CREDIT_NOTE,
    "credit_note_missing": MatchStatus.MISSING_IN_ACCOUNTING,
    "payment_matched": MatchStatus.MATCHED,
    "payment_run_grouped": MatchStatus.MATCHED,
    "payment_amount_date": MatchStatus.MATCHED_WITH_TOLERANCE,
}


@pytest.fixture(params=["statement_pdf", "statement_xlsx", "statement_csv"])
def reconciled(request, database, tmp_path):
    data = demo_data.generate(40)
    paths = demo_data.write_files(data, tmp_path / "demo")
    service = ReconciliationService(database)
    importer = ImportService(database, replace(get_settings(), upload_dir=tmp_path / "uploads"))
    recon_id = service.create(
        NewReconciliation(
            supplier_name=demo_data.DEMO_SUPPLIER,
            supplier_account=demo_data.DEMO_SUPPLIER_ACCOUNT,
            currency="GBP",
            period_start=demo_data.PERIOD_START,
            period_end=demo_data.PERIOD_END,
            is_demo=True,
        )
    )
    for role, key in ((FileRole.STATEMENT, request.param),
                      (FileRole.ACCOUNTING, "accounting_csv"),
                      (FileRole.PAYMENTS, "payments_csv")):  # fmt: skip
        outcome = importer.upload(recon_id, role, paths[key].name, paths[key].read_bytes())
        assert outcome.mapping_complete, (role, outcome.missing_required)
        assert outcome.result.report.errors == 0, outcome.result.report.issues
    run = service.run(recon_id)
    return data, service, recon_id, run, database


def _rows_by_reference(service, recon_id, run):
    lookup = service.load_data(recon_id).lookup()
    rows: dict[str, list] = {}
    for row in build_rows(run.results, lookup):
        rows.setdefault(row.reference, []).append(row)
        if row.statement_reference:
            rows.setdefault(row.statement_reference, []).append(row)
    return rows


def test_every_demo_scenario_is_detected(reconciled):
    data, service, recon_id, run, _ = reconciled
    rows = _rows_by_reference(service, recon_id, run)

    for scenario, expected in EXPECTED_STATUS.items():
        for reference in data.manifest[scenario]:
            statuses = {r.status for r in rows[reference]}
            assert expected in statuses, (scenario, reference, statuses)

    for reference in data.manifest["duplicate_in_accounting"]:
        assert {r.status for r in rows[reference]} == {MatchStatus.MATCHED, MatchStatus.DUPLICATE}
    for key in ("accounting_only", "accounting_only_after_period"):
        for reference in data.manifest[key]:
            assert rows[reference][0].status == MatchStatus.NOT_ON_SUPPLIER_STATEMENT
    timing = rows[data.manifest["accounting_only_after_period"][0]][0]
    assert "timing difference" in timing.reason
    assert rows["PMT-24904"][0].status == MatchStatus.PAYMENT_MISMATCH

    unresolved = [r for r in run.results if r.status == MatchStatus.UNRESOLVED]
    assert len(unresolved) == 1 and len(unresolved[0].related_txn_ids) == 2

    fuzzy = next(r for r in run.results if r.matching_method == MatchMethod.FUZZY_REFERENCE)
    assert fuzzy.review_status == ReviewStatus.PENDING  # never silently accepted


def test_summary_and_balance(reconciled):
    data, _, _, run, _ = reconciled
    summary, balance = run.summary, run.balance
    assert summary.statement_transactions == len(data.statement)
    assert summary.accounting_transactions == len(
        [b for b in data.bills if b.supplier == demo_data.DEMO_SUPPLIER]
    )  # the other supplier's rows were excluded
    assert balance.status == BalanceStatus.BALANCED
    assert balance.opening_balance == data.opening_balance
    assert balance.supplier_closing_balance == data.closing_balance
    assert balance.unexplained_difference == Decimal("0.00")
    assert sum(balance.explained_difference.values()) == balance.statement_vs_accounting_difference
    assert summary.suggested_outcome == ReconciliationOutcome.NOT_RECONCILED  # reviews pending


def test_review_then_export_workbook_and_report(reconciled, tmp_path):
    _, service, recon_id, run, database = reconciled
    reviews = ReviewService(database, service)
    for result in run.results:
        if result.review_status == ReviewStatus.PENDING:
            reviews.review(recon_id, result.result_id, ReviewAction.ACCEPT, "Checked")
    stored = service.stored_run(recon_id)
    assert stored.summary.pending_reviews == 0
    assert stored.summary.suggested_outcome == ReconciliationOutcome.RECONCILED_WITH_EXCEPTIONS
    service.finalise(recon_id, ReconciliationOutcome.RECONCILED_WITH_EXCEPTIONS, "Sent to client")

    context = build_export_context(service, recon_id)
    workbook = load_workbook(io.BytesIO(ExportService().workbook_bytes(context)))
    for sheet in ("Summary", "Matched", "Missing", "Duplicates", "Amount Mismatches",
                  "Date Mismatches", "Reference Mismatches", "Credit Notes", "Payment Issues",
                  "Accounting Only", "Unresolved", "Balance Check"):  # fmt: skip
        assert sheet in workbook.sheetnames
    summary = {r[0].value: r[1].value for r in workbook["Summary"].iter_rows(min_row=4)
               if r[0].value}  # fmt: skip
    assert summary["Overall status"] == "Reconciled With Exceptions"
    assert summary["Missing in accounting"] == 4
    assert summary["Duplicates"] == 2
    assert workbook["Missing"].max_row == 5  # header + 4 rows
    assert workbook["Amount Mismatches"].max_row == 3
    assert "DEMO DATA" in workbook["Summary"]["A2"].value

    pdf = ReportService().pdf_bytes(context)
    assert pdf.startswith(b"%PDF") and len(pdf) > 5_000


# ----------------------------------------------------------------- UI journey


@pytest.fixture
def app_env(tmp_path, monkeypatch):
    monkeypatch.setenv("SSR_DATA_DIR", str(tmp_path / "data"))
    config.get_settings.cache_clear()
    database_module.get_database.cache_clear()
    st.cache_resource.clear()
    yield
    config.get_settings.cache_clear()
    database_module.get_database.cache_clear()
    st.cache_resource.clear()


def test_user_journey_through_the_interface(app_env):
    app = AppTest.from_file(str(PROJECT_ROOT / "app.py"), default_timeout=60)
    app.run()
    assert not app.exception

    # Dashboard: create the demo reconciliation.
    demo = next(b for b in app.button if b.label == "Load demo reconciliation")
    demo.click().run()
    assert not app.exception
    recon_id = app.session_state[CURRENT_KEY]

    # Reconciliation page: run it.
    app.switch_page("pages/reconciliation.py").run()
    run_button = next(b for b in app.button if b.label == "Run reconciliation")
    run_button.click().run()
    assert not app.exception
    text = " ".join(m.value for m in app.markdown)
    assert "missing from accounting" in text

    # Remaining pages render with results.
    for page in ("pages/exceptions.py", "pages/reports.py", "pages/dashboard.py"):
        app.switch_page(page).run()
        assert not app.exception, page

    service = ReconciliationService(database_module.get_database())
    stored = service.stored_run(recon_id)
    assert stored is not None and stored.summary.count(MatchStatus.DUPLICATE) == 2
